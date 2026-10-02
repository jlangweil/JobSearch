from __future__ import annotations

import io
import shutil
import time

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from jobhunt.config import PROJECT_ROOT
from jobhunt.db import queries
from jobhunt.facts import bank
from jobhunt.sources.normalize import normalize_raw
from jobhunt.tailor import render
from jobhunt.web.app import create_app
from tests.conftest import FakeLLM, FakeSource, raw_job
from tests.test_application import FABRICATED, ScriptedLLM


def seed(engine, n: int = 6) -> list[int]:  # noqa: ANN001
    ids = []
    for i in range(n):
        remote = "Remote (US)" if i % 2 == 0 else "Parsippany, NJ"
        rec = normalize_raw(raw_job(i, title=f"Integration Architect {i}", company=f"Co {i}", location=remote,
                                    description="<p>Remote-first team. Own the API platform.</p>" if i % 2 == 0
                                    else "<p>Hybrid in Parsippany. Own the API platform.</p>"))
        _, job_id = queries.upsert_job(engine, rec)
        total = 60 + i * 6  # 60, 66, 72, 78, 84, 90
        queries.record_score(engine, job_id, {"total": total, "skill_match": total, "role_fit": total,
                                              "domain_fit": total,
                                              "seniority_fit": total, "comp_fit": 50, "logistics_fit": 80, "focus_fit": 50,
                                              "ghost_score": 10, "rationale": f"rationale {i}"})
        ids.append(job_id)
    return ids


@pytest.fixture
def client(cfg, engine, hashing_embedder):
    app = create_app(cfg, start_scheduler=False,
                     runtime_factories={"llm_factory": FakeLLM, "embedder_factory": lambda: hashing_embedder,
                                        "sources_factory": lambda names: []})
    with TestClient(app) as c:
        yield c


def test_root_and_default_queue_redirects(client):
    r = client.get("/", follow_redirects=False)
    assert r.status_code == 303 and r.headers["location"] == "/home"
    r = client.get("/queue", follow_redirects=False)
    assert r.status_code == 303 and r.headers["location"] == "/queue?status=scored&min_score=70&sort=total"


def test_queue_filters_live_in_the_url_and_survive_restart(cfg, engine, client, hashing_embedder):
    seed(engine)
    url = "/queue?min_score=80&remote=remote"
    r = client.get(url)
    assert r.status_code == 200
    assert "Integration Architect 4" in r.text  # 84, remote
    assert "Integration Architect 5" not in r.text  # 90 but hybrid
    assert "Integration Architect 2" not in r.text  # remote but 72
    assert 'value="80"' in r.text  # the filter form reflects the URL
    # a "restart": a brand-new app instance on the same database renders the same bookmark
    app2 = create_app(cfg, start_scheduler=False, runtime_factories={"embedder_factory": lambda: hashing_embedder})
    with TestClient(app2) as c2:
        r2 = c2.get(url)
    assert r2.status_code == 200 and "Integration Architect 4" in r2.text and "Integration Architect 5" not in r2.text


def test_export_matches_filtered_view(engine, client):
    seed(engine)
    r = client.get("/queue/export?status=scored&min_score=70&sort=total")
    assert r.status_code == 200
    ws = load_workbook(io.BytesIO(r.content)).active
    titles = [row[4] for row in ws.iter_rows(min_row=2, values_only=True)]
    assert titles == ["Integration Architect 5", "Integration Architect 4", "Integration Architect 3",
                      "Integration Architect 2"]


def test_triage_via_htmx_and_plain_forms_lands_in_decisions(engine, client):
    ids = seed(engine)
    r = client.post(f"/jobs/{ids[5]}/reject", data={"reason": "not integration-heavy"}, headers={"HX-Request": "true"})
    assert r.status_code == 200 and f'id="job-{ids[5]}"' in r.text and "rejected" in r.text
    r = client.post(f"/jobs/{ids[4]}/reject", data={"reason": ""}, headers={"HX-Request": "true"})
    assert "reason is required" in r.text
    r = client.post(f"/jobs/{ids[4]}/snooze", data={"until": "2030-01-01", "next": "/queue?status=scored"},
                    follow_redirects=False)
    assert r.status_code == 303 and r.headers["location"] == "/queue?status=scored"
    actions = {d["job_id"]: d["action"] for d in
               [*queries.decisions_for_job(engine, ids[5]), *queries.decisions_for_job(engine, ids[4])]}
    assert actions == {ids[5]: "reject", ids[4]: "snooze"}
    assert queries.get_job(engine, ids[4])["snooze_until"] == "2030-01-01T00:00:00Z"
    detail = client.get(f"/jobs/{ids[5]}")
    assert detail.status_code == 200 and "not integration-heavy" in detail.text


def test_run_from_browser_with_live_progress_while_queue_stays_responsive(cfg, engine, hashing_embedder):
    seed(engine)

    class SlowSource(FakeSource):
        def fetch(self, since):  # noqa: ANN001
            for job in self.jobs:
                time.sleep(0.05)
                yield job

    jobs = [raw_job(100 + i, title=f"Principal API Architect {i}", company=f"Slow Co {i}") for i in range(30)]
    app = create_app(cfg, start_scheduler=False, runtime_factories={
        "llm_factory": FakeLLM, "embedder_factory": lambda: hashing_embedder,
        "sources_factory": lambda names: [SlowSource("workday", jobs)]})
    with TestClient(app) as c:
        r = c.post("/runs", data={"max_to_score": "50"}, follow_redirects=False)
        assert r.status_code == 303
        run_path = r.headers["location"]
        run_id = int(run_path.rsplit("/", 1)[1])
        # a second start is refused by the lock while the first is running
        r2 = c.post("/runs", data={}, follow_redirects=False)
        assert "already+in+progress" in r2.headers["location"] or "already%20in%20progress" in r2.headers["location"]
        saw_running, deadline = False, time.time() + 60
        while time.time() < deadline:
            t0 = time.perf_counter()
            q = c.get("/queue?status=scored&min_score=0&sort=total")
            assert q.status_code == 200 and "database is locked" not in q.text
            assert time.perf_counter() - t0 < 2.0
            p = c.get(f"{run_path}/progress", headers={"HX-Request": "true"})
            assert p.status_code == 200
            if 'hx-trigger="every 2s"' in p.text:
                saw_running = True
            run = queries.get_run(engine, run_id)
            if run["status"] != "running":
                break
            time.sleep(0.1)
        assert saw_running
        assert run["status"] == "complete", run["errors"]
        assert run["counts"]["scored"] == 30
        page = c.get(run_path)
        assert "✓" in page.text and 'hx-trigger="every 2s"' not in page.text


def test_search_settings_backup_facts_and_applications_pages(cfg, engine, client):
    seed(engine)
    r = client.get("/search?q=Parsippany")
    assert r.status_code == 200 and "Integration Architect 1" in r.text
    r = client.get("/settings")
    assert r.status_code == 200 and "Scoring weights" in r.text
    form = {"skill_match": "1", "role_fit": "0", "domain_fit": "0", "seniority_fit": "0", "comp_fit": "0",
            "logistics_fit": "0", "focus_fit": "0"}
    r = client.post("/settings/weights", data=form, follow_redirects=False)
    assert r.status_code == 303 and "recomputed" in r.headers["location"]
    with open(cfg.config_path, encoding="utf-8") as fh:
        assert "weights: {skill_match: 1.0, role_fit: 0.0" in fh.read()
    s = queries.latest_score(engine, queries.list_queue(engine, queries.QueueFilters(status="any"))[0][0]["id"])
    assert s["total"] == pytest.approx(s["skill_match"] * (1 - cfg.scoring.ghost_penalty * s["ghost_score"] / 100),
                                       abs=0.1)
    r = client.post("/settings/backup", follow_redirects=False)
    assert r.status_code == 303 and "backup+written" in r.headers["location"]
    assert list(cfg.paths.exports.glob("jobhunt-backup-*.db"))
    shutil.copy(PROJECT_ROOT / "data" / "facts.example.yaml", cfg.paths.facts)
    r = client.post("/facts", data={"id": "ach-new", "kind": "achievement", "parent_id": "role-example-current",
                                    "text": "Automated 40 deployments.", "tags": "ci, automation",
                                    "metrics": "deployments: 40"}, follow_redirects=False)
    assert r.status_code == 303 and "added" in r.headers["location"], r.headers["location"]
    with open(cfg.paths.facts, encoding="utf-8") as fh:
        assert "ach-new" in fh.read()
    assert "ach-new" in client.get("/facts?q=deployments").text
    bad = client.post("/facts", data={"id": "ach-bad", "kind": "achievement", "parent_id": "nope", "text": "x"},
                      follow_redirects=False)
    assert "error=" in bad.headers["location"]
    with open(cfg.paths.facts, encoding="utf-8") as fh:
        assert "ach-bad" not in fh.read()  # the file is restored when validation fails
    assert client.get("/applications").status_code == 200
    assert client.get("/runs").status_code == 200


def test_workflow_pages_and_hand_edit_validation_in_the_browser(cfg, engine, hashing_embedder):
    shutil.copy(PROJECT_ROOT / "data" / "facts.example.yaml", cfg.paths.facts)
    render.make_default_template(cfg.paths.resume_template)
    bank.sync(engine, cfg, hashing_embedder)
    ids = seed(engine, 1)
    llm = ScriptedLLM()
    app = create_app(cfg, start_scheduler=False,
                     runtime_factories={"llm_factory": lambda: llm, "embedder_factory": lambda: hashing_embedder})
    with TestClient(app) as c:
        r = c.post(f"/jobs/{ids[0]}/shortlist", data={"reason": "fit"}, follow_redirects=False)
        assert r.headers["location"] == f"/jobs/{ids[0]}/workflow"
        for _ in range(100):
            page = c.get(f"/jobs/{ids[0]}/workflow")
            if "Resume draft" in page.text:
                break
            time.sleep(0.1)
        assert "Resume draft" in page.text and "{fact:ach-example-gateway}" in page.text
        assert "ATS keyword match" in page.text
        from jobhunt.services import workflow as wf
        from jobhunt.web.common import rt

        editor = wf.view(rt(), ids[0])["editor"]
        c.post(f"/jobs/{ids[0]}/workflow/resume",
               data={"action": "edit", "kind": "resume", "editor": editor.replace("time 60%", "time 85%")})
        for _ in range(100):
            page = c.get(f"/jobs/{ids[0]}/workflow")
            if "Validation failed" in page.text:
                break
            time.sleep(0.1)
        assert "Validation failed" in page.text and "85%" in page.text and "unsupported_number" in page.text
        assert FABRICATED  # fixture import sanity
        r = c.post(f"/jobs/{ids[0]}/workflow/resume", data={"action": "confirm"}, follow_redirects=False)
        assert "not+a+valid+action" in r.headers["location"]


def test_basic_auth_when_password_set(cfg, hashing_embedder):
    cfg.secrets.WEB_PASSWORD = "s3cret"
    app = create_app(cfg, start_scheduler=False, runtime_factories={"embedder_factory": lambda: hashing_embedder})
    with TestClient(app) as c:
        assert c.get("/queue?status=scored").status_code == 401
        assert c.get("/queue?status=scored", auth=("me", "wrong")).status_code == 401
        assert c.get("/queue?status=scored", auth=("me", "s3cret")).status_code == 200
    cfg.secrets.WEB_PASSWORD = None
