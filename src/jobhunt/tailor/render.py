"""Rendering: Windows-safe output paths, ATS-compliant .docx from validated markdown, optional PDF."""

from __future__ import annotations

import re
import subprocess
import unicodedata
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

from jobhunt.sources.normalize import slugify

_ILLEGAL = re.compile(r'[<>:"/\\|?*\x00-\x1f]')
_RESERVED = {"CON", "PRN", "AUX", "NUL", *(f"COM{i}" for i in range(1, 10)), *(f"LPT{i}" for i in range(1, 10))}
COMPONENT_MAX = 64

ATS_SAFE_FONTS = {"calibri", "arial", "helvetica", "georgia", "garamond", "cambria", "times new roman", "verdana",
                  "tahoma", "aptos", "book antiqua", "trebuchet ms"}


class RenderError(RuntimeError):
    pass


# --------------------------------------------------------------------------------------------- path sanitizer


def sanitize_component(name: str, max_len: int = COMPONENT_MAX) -> str:
    """One path component that Windows will accept: no illegal/control chars, no trailing dot/space,
    no reserved device name, at most max_len characters."""
    s = unicodedata.normalize("NFC", name or "")
    s = _ILLEGAL.sub("-", s)
    s = re.sub(r"\s+", " ", s)
    s = re.sub(r"-{2,}", "-", s).strip(" -")
    s = s[:max_len]
    s = s.rstrip(". ")
    if not s:
        s = "untitled"
    stem = s.split(".")[0].strip().upper()
    if stem in _RESERVED:
        s = f"_{s}"[:max_len].rstrip(". ")
    return s


def output_dir(root: Path, company: str, title: str, job_id: int, *, max_path_chars: int = 240,
               reserve_for_filename: int = 48) -> Path:
    """output/<company-slug>-<title>-<jobid>/, guaranteed to leave room for a filename under max_path_chars."""
    root = Path(root).resolve()
    suffix = f"-{job_id}"
    prefix = slugify(company)[:24]
    budget = min(COMPONENT_MAX, max_path_chars - len(str(root)) - 1 - reserve_for_filename)
    if budget < len(suffix) + 8:
        raise RenderError(f"output_root {root} is too deep to stay under {max_path_chars} characters")
    title_part = sanitize_component(title, COMPONENT_MAX)
    room = budget - len(prefix) - len(suffix) - 1
    name = sanitize_component(f"{prefix}-{title_part[:max(room, 1)].rstrip('. -')}", budget - len(suffix)) + suffix
    name = sanitize_component(name, budget)
    return root / name


def document_filename(person: str | None, kind: str, ext: str, version: int | None = None) -> str:
    """The current document keeps the plain name you upload; an older version gets a -v<N> suffix so re-rendering
    one never overwrites the file you are applying with."""
    label = "Resume" if kind == "resume" else "Cover-Letter"
    if version:
        label = f"{label}-v{version}"
    base = "-".join(p for p in re.split(r"\s+", (person or "").strip()) if p)
    return sanitize_component(f"{base}-{label}.{ext}" if base else f"{label}.{ext}")


# --------------------------------------------------------------------------------------------- docx


def _clear_body(doc) -> None:  # noqa: ANN001
    body = doc.element.body
    for child in list(body):
        if child.tag.endswith("}sectPr"):
            continue
        body.remove(child)


def _style(doc, *names: str):  # noqa: ANN001, ANN202
    for n in names:
        try:
            return doc.styles[n]
        except KeyError:
            continue
    return None


def _add(doc, text: str, *names: str, bold: bool = False, size: float | None = None,  # noqa: ANN001
         align_center: bool = False):  # noqa: ANN202
    style = _style(doc, *names)
    p = doc.add_paragraph(style=style) if style is not None else doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold or None
    if size:
        run.font.size = Pt(size)
    if align_center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return p


def markdown_to_docx(markdown: str, out_path: Path, *, template: Path | None, name: str, contact_line: str,
                     kind: str, role_layout: str = "separate_lines", headline: str | None = None) -> Path:
    """Single-column, table-free, image-free document using built-in Heading / List Bullet styles so ATS
    parsers read it in order. Contact details live in the body, never in headers/footers."""
    doc = Document(str(template)) if template and Path(template).exists() else Document()
    if missing := missing_styles(doc):
        raise RenderError(f"resume template {Path(template).name if template else ''} is missing Word styles "
                          f"{', '.join(missing)} — headings or bullets would lose their formatting. Use a styles-only "
                          "template (see README → Resume styling), not a copy of a finished resume.")
    _clear_body(doc)
    for doc_section in doc.sections:
        for part in (doc_section.header, doc_section.footer):
            for p in part.paragraphs:
                for r in p.runs:
                    r.text = ""
    if name:
        _add(doc, name, "Title")
    if contact_line:
        _add(doc, contact_line, "Normal", align_center=True)
    if headline and kind == "resume":  # ATS Super Match: the role you are applying as, under your contact line
        _add(doc, headline, "Normal", bold=True, align_center=True)
    section = "other"  # which resume section the role headings below belong to
    for raw in (markdown or "").replace("\r\n", "\n").split("\n"):
        line = raw.rstrip()
        if not line.strip():
            if kind == "cover_letter":
                continue
            continue
        if line.startswith("## "):
            section = _section_of(line[3:].strip())
            label = line[3:].strip()
            if kind == "resume":
                # the section labels Workday's extractor looks for
                label = {"experience": "WORK EXPERIENCE"}.get(label.lower(), label.upper())
            _add(doc, label, "Heading 1")
        elif line.startswith("### "):
            _role_heading(doc, line[4:].strip(), section, role_layout)
        elif re.match(r"^\s*[-*]\s+", line):
            _add(doc, re.sub(r"^\s*[-*]\s+", "", line), "List Bullet")
        elif section in ("skills", "projects") and kind == "resume" and re.match(r"^[^:,]{1,40}:\s*\S", line.strip()):
            _skills_line(doc, line.strip())
        else:
            _add(doc, line.strip(), "Normal")
    props = doc.core_properties
    props.title = f"{name} {'Resume' if kind == 'resume' else 'Cover Letter'}".strip()
    props.author = name or ""
    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    try:
        doc.save(str(out_path))
    except PermissionError as exc:
        raise RenderError(f"{out_path.name} is open in another program (probably Word). Close it and retry.") from exc
    return out_path


def _section_of(heading: str) -> str:
    low = heading.strip().lower()
    if low.startswith("skills"):
        return "skills"
    if low.startswith("projects"):
        return "projects"
    return "experience" if low.startswith("experience") else ("education" if low.startswith("education") else "other")


def _skills_line(doc, line: str) -> None:  # noqa: ANN001
    """One Skills group: a bold "Label:" then the skills, comma-separated, with a little space after. Still one plain
    paragraph of text (bold is just a run property), so parsers read "Integration and APIs: REST, MuleSoft, ..."
    exactly as before; only people see the label stand out."""
    label, _, rest = line.partition(":")
    p = doc.add_paragraph(style=_style(doc, "Normal"))
    p.add_run(f"{label.strip()}:").bold = True
    p.add_run(" " + ", ".join(x.strip() for x in rest.split(",") if x.strip()))
    p.paragraph_format.space_after = Pt(3)


def _role_heading(doc, heading: str, section: str, layout: str) -> None:  # noqa: ANN001
    """Write one role (or degree) heading as plain text lines, which is what resume parsers read.

    Workday and friends segment work history from the normal text flow: employer on its own line, then the title,
    then the dates. Two things broke this before — the whole role on one line (a dash inside a title, e.g.
    "Solution Lead – Enterprise Integration", made title/employer/dates unguessable), and putting the employer in a
    Word *heading* style, which a parser treats as a section boundary and skips (Workday then took "Microsoft" out
    of a bullet as the employer). So: Normal style with bold for emphasis, and MM/YYYY dates."""
    parts = [p.strip() for p in heading.split("|") if p.strip()]
    if layout != "separate_lines" or section not in ("experience", "education") or len(parts) < 2:
        _add(doc, heading.replace(" | ", "  |  "), "Heading 2")
        return
    title, org = parts[0], parts[1]
    org_line = _add(doc, org, "Normal", bold=True)
    org_line.paragraph_format.space_before = Pt(8)
    _add(doc, _plain_title(title), "Normal")
    if dates := _role_dates(" ".join(parts[2:])):
        _add(doc, dates, "Normal")


def _plain_title(title: str) -> str:
    """A dash inside a title invites a parser to split it into "title - employer"; a comma reads the same to a
    human and doesn't."""
    return re.sub(r"\s+[–—-]\s+", ", ", title)


_MONTHS = {m: i for i, m in enumerate(
    ["jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"], start=1)}
SPLIT_DATES = re.compile(r"\s*(?:[–—-]|to|until)\s*", re.IGNORECASE)
_PRESENT = re.compile(r"^(present|current|now|to date)$", re.IGNORECASE)


def _one_date(value: str) -> str | None:
    v = value.strip(" ,")
    if not v:
        return None
    if _PRESENT.match(v):
        return "Present"
    if m := re.fullmatch(r"(\d{1,2})[/-](\d{4})", v):
        return f"{int(m.group(1)):02d}/{m.group(2)}"
    if m := re.fullmatch(r"([A-Za-z]{3,9})\.?\s+(\d{4})", v):
        month = _MONTHS.get(m.group(1)[:3].lower())
        return f"{month:02d}/{m.group(2)}" if month else v
    if re.fullmatch(r"(19|20)\d{2}", v):
        return v
    return v


def _role_dates(text: str) -> str:
    """"Sep 2018 - Oct 2025" -> "09/2018 - 10/2025": the date format parsers map most reliably."""
    raw = text.strip()
    if not raw:
        return ""
    parts = re.split(SPLIT_DATES, raw, maxsplit=1)
    ends = [d for d in (_one_date(p) for p in parts) if d]
    return " - ".join(ends) if ends else raw


def check_docx_ats(path: Path) -> list[str]:
    """Structural ATS checks on a rendered .docx."""
    problems: list[str] = []
    doc = Document(str(path))
    if doc.tables:
        problems.append("contains tables")
    if doc.inline_shapes:
        problems.append("contains images/shapes")
    with zipfile.ZipFile(path) as z:
        xml = z.read("word/document.xml").decode("utf-8", errors="replace")
        if "w:txbxContent" in xml:
            problems.append("contains text boxes")
        if re.search(r'<w:cols [^>]*w:num="([2-9])"', xml):
            problems.append("uses multiple columns")
    for section in doc.sections:
        for label, part in (("header", section.header), ("footer", section.footer)):
            if any(p.text.strip() for p in part.paragraphs):
                problems.append(f"text in page {label} (ATS parsers often skip it)")
    fonts = {s.font.name.lower() for s in doc.styles if getattr(s, "font", None) is not None and s.font.name}
    if unsafe := sorted(fonts - ATS_SAFE_FONTS):
        problems.append(f"non-standard fonts in styles: {', '.join(unsafe)}")
    return problems


def render_pdf(docx_path: Path, soffice: Path | None, timeout: int = 120) -> Path | None:
    """Convert with LibreOffice when configured; returns None (not an error) when it is unavailable."""
    if soffice is None or not Path(soffice).exists():
        return None
    out_dir = Path(docx_path).parent
    try:
        subprocess.run(
            [str(soffice), "--headless", "--convert-to", "pdf", "--outdir", str(out_dir), str(docx_path)],
            check=True, capture_output=True, encoding="utf-8", errors="replace", timeout=timeout,
        )
    except PermissionError as exc:
        raise RenderError(f"cannot write PDF next to {Path(docx_path).name}: close it in any viewer and retry") from exc
    except (subprocess.CalledProcessError, subprocess.TimeoutExpired, OSError) as exc:
        raise RenderError(f"LibreOffice PDF conversion failed: {exc}") from exc
    pdf = Path(docx_path).with_suffix(".pdf")
    return pdf if pdf.exists() else None


REQUIRED_STYLES = ("Normal", "Title", "Heading 1", "Heading 2", "List Bullet")


def missing_styles(doc) -> list[str]:  # noqa: ANN001
    return [name for name in REQUIRED_STYLES if _style(doc, name) is None]


def _force_font(style, name: str) -> None:  # noqa: ANN001
    """Set a style's font explicitly, removing theme-font references that Word would otherwise prefer."""
    from docx.oxml.ns import qn

    style.font.name = name
    rpr = style.element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is not None:
        for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            fonts.attrib.pop(qn(attr), None)
        for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            fonts.set(qn(attr), name)


def _black(style) -> None:  # noqa: ANN001
    from docx.oxml.ns import qn
    from docx.shared import RGBColor

    style.font.color.rgb = RGBColor(0, 0, 0)
    color = style.element.get_or_add_rPr().find(qn("w:color"))
    if color is not None:
        for attr in ("w:themeColor", "w:themeShade", "w:themeTint"):
            color.attrib.pop(qn(attr), None)


def make_default_template(path: Path, *, font: str = "Calibri", body_pt: float = 11, name_pt: float = 14,
                          section_pt: float = 12, role_pt: float = 11, margin_in: float = 0.5,
                          center_sections: bool = True) -> Path:
    """Create an ATS-friendly style template: styles only, no body text, nothing in headers/footers.

    Defaults mirror a compact, conventional resume: one font, black text, bold all-caps section headings
    (the renderer upper-cases them), bold role lines, tight spacing, 0.5" margins."""
    from docx.oxml.ns import qn
    from docx.shared import Inches

    doc = Document()
    spec = {
        "Normal": dict(size=body_pt, bold=False, before=0, after=0),
        "Title": dict(size=name_pt, bold=True, before=0, after=0),
        "Heading 1": dict(size=section_pt, bold=True, before=8, after=2),
        "Heading 2": dict(size=role_pt, bold=True, before=6, after=0),
        "List Bullet": dict(size=body_pt, bold=False, before=0, after=0),
    }
    for style_name, v in spec.items():
        st = doc.styles[style_name]
        _force_font(st, font)
        _black(st)
        st.font.size = Pt(v["size"])
        st.font.bold = v["bold"]
        st.font.italic = False
        pf = st.paragraph_format
        pf.space_before, pf.space_after = Pt(v["before"]), Pt(v["after"])
        pf.line_spacing = 1.0
        ppr = st.element.get_or_add_pPr()
        border = ppr.find(qn("w:pBdr"))
        if border is not None:  # Word's built-in Title has a bottom rule; keep the page plain
            ppr.remove(border)
    if center_sections:
        doc.styles["Heading 1"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.styles["Title"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for s in doc.styles:
        if getattr(s, "font", None) is not None and s.font.name and s.font.name.lower() not in ATS_SAFE_FONTS:
            s.font.name = font
    for section in doc.sections:
        section.left_margin = section.right_margin = Inches(margin_in)
        section.top_margin = section.bottom_margin = Inches(margin_in)
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))
    return path
