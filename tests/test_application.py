from __future__ import annotations

import shutil
from pathlib import Path

import pytest
from docx import Document

from jobhunt.config import PROJECT_ROOT
from jobhunt.db import queries
from jobhunt.facts import bank
from jobhunt.graphs.runtime import init_runtime
from jobhunt.llm.schemas import (
    BulletFactRef,
    CoverDraft,
    ParagraphFactRef,
    Requirement,
    RequirementList,
    ResumeDraft,
)
from jobhunt.services import triage, workflow
from jobhunt.sources.normalize import normalize_raw
from jobhunt.tailor import render
from tests.conftest import raw_job

TITLE = "Sr. Architect: API/Integration (Remote) "

GOOD_RESUME = """## Summary
PLACEHOLDER: Integration architect with 25 years of experience leading API and iPaaS platforms in insurance and financial services.

## Skills
Integration: MuleSoft, REST APIs, API gateways, Kafka, AWS

## Experience
### Lead Integration Architect | Example Insurance Co | Mar 2019 – Present
- Migrated 120 partner integrations to a centralized API gateway, cutting partner onboarding time 60%.
- Managed a $2.5M annual integration platform budget across 3 delivery teams.

### Senior Integration Engineer | Example Health Plan | Jun 2012 – Feb 2019
- Introduced MuleSoft as the integration standard, reaching 99.95% uptime for eligibility APIs.

## Education
### B.S. Computer Science | Example University | 1999

## Certifications
- MuleSoft Certified Integration Architect"""
GOOD_MAP = [BulletFactRef(bullet_index=0, fact_id="ach-example-gateway"),
            BulletFactRef(bullet_index=1, fact_id="ach-example-budget"),
            BulletFactRef(bullet_index=2, fact_id="ach-example-mulesoft")]
FABRICATED = GOOD_RESUME.replace("onboarding time 60%", "onboarding time 85%")
# a made-up number in the summary: no mechanical fix, so it goes back to the model
FABRICATED_SUMMARY = GOOD_RESUME.replace("with 25 years of experience", "with 31 years of experience")
COVER = """Dear Hiring Manager,

I am excited to apply for this role.

At Example Insurance Co I migrated 120 partner integrations to a centralized API gateway, cutting partner onboarding time 60%.

Sincerely,"""


class ScriptedLLM:
    def __init__(self) -> None:
        self.resumes: list[str] = [GOOD_RESUME]
        self.calls: dict[str, int] = {}

    def structured(self, schema, *, system, user, purpose, model, prompt_version, job_id=None, max_tokens=4096,
                   cache_system=False):  # noqa: ANN001
        self.calls[purpose] = self.calls.get(purpose, 0) + 1
        if schema is RequirementList:
            return RequirementList(requirements=[
                Requirement(text="API platform ownership", kind="must", keywords=["api gateway", "mulesoft"]),
                Requirement(text="Kafka", kind="nice", keywords=["kafka"]),
                *([Requirement(text="Program delivery", kind="must", keywords=list(self.requirements_keywords))]
                  if getattr(self, "requirements_keywords", None) else []),
            ])
        if schema is ResumeDraft:
            md = self.resumes.pop(0) if len(self.resumes) > 1 else self.resumes[0]
            return ResumeDraft(markdown=md, fact_map=GOOD_MAP)
        if schema is CoverDraft:
            return CoverDraft(markdown=COVER, fact_map=[ParagraphFactRef(paragraph_index=1, fact_id="ach-example-gateway")])
        raise AssertionError(schema)


# --------------------------------------------------------------------------------------------- filenames / docx


def test_sanitizer_handles_real_shaped_title(tmp_path):
    comp = render.sanitize_component(TITLE)
    assert not any(c in comp for c in '<>:"/\\|?*')
    assert not comp.endswith((" ", "."))
    assert render.sanitize_component("CON") == "_CON"
    assert render.sanitize_component("lpt1.txt") == "_lpt1.txt"
    assert render.sanitize_component("nul ") == "_nul"
    assert len(render.sanitize_component("x" * 300)) == 64
    with pytest.raises(render.RenderError, match="too deep"):
        render.output_dir(tmp_path / ("d" * 100) / ("e" * 100), "Prudential Financial", TITLE, 1)
    deep = tmp_path / ("d" * 40)
    out = render.output_dir(deep, "Prudential Financial", TITLE, 12345)
    assert out.name.endswith("-12345") and len(out.name) <= 64
    full = out / render.document_filename("Jane Q. Public", "resume", "docx")
    assert len(str(full)) < 240
    full.parent.mkdir(parents=True)
    full.write_text("ok", encoding="utf-8")  # Windows accepts the path
    assert full.exists()


def test_rendered_docx_is_ats_clean(tmp_path):
    template = render.make_default_template(tmp_path / "template.docx")
    out = render.markdown_to_docx(GOOD_RESUME, tmp_path / "Jane-Resume.docx", template=template, name="Jane Public",
                                  contact_line="Newark, NJ | jane@example.com", kind="resume")
    assert render.check_docx_ats(out) == []
    doc = Document(str(out))
    texts = [p.text for p in doc.paragraphs]
    assert texts[0] == "Jane Public" and "jane@example.com" in texts[1]
    assert "WORK EXPERIENCE" in texts  # the label Workday's extractor looks for
    styles = {p.style.name for p in doc.paragraphs}
    assert {"Heading 1", "List Bullet"} <= styles


def test_docx_with_table_fails_ats_check(tmp_path):
    doc = Document()
    doc.add_table(rows=1, cols=2)
    doc.sections[0].header.paragraphs[0].text = "Jane Public"
    path = tmp_path / "bad.docx"
    doc.save(str(path))
    problems = render.check_docx_ats(path)
    assert any("tables" in p for p in problems) and any("header" in p for p in problems)


def test_word_file_lock_becomes_clear_error(tmp_path, monkeypatch):
    def locked(self, path):  # noqa: ANN001
        raise PermissionError(13, "Permission denied")

    monkeypatch.setattr("docx.document.Document.save", locked)
    with pytest.raises(render.RenderError, match="open in another program"):
        render.markdown_to_docx(GOOD_RESUME, tmp_path / "x.docx", template=None, name="J", contact_line="",
                                kind="resume")


# --------------------------------------------------------------------------------------------- workflow


@pytest.fixture
def app_rt(cfg, engine, hashing_embedder):
    shutil.copy(PROJECT_ROOT / "data" / "facts.example.yaml", cfg.paths.facts)
    render.make_default_template(cfg.paths.resume_template)
    llm = ScriptedLLM()
    rt = init_runtime(cfg, llm_factory=lambda: llm, embedder_factory=lambda: hashing_embedder)
    bank.sync(engine, cfg, hashing_embedder)
    rec = normalize_raw(raw_job(1, title=TITLE, company="Prudential"))
    _, job_id = queries.upsert_job(engine, rec)
    queries.record_score(engine, job_id, {"total": 82.0, "skill_match": 80, "domain_fit": 80, "seniority_fit": 80,
                                          "comp_fit": 80, "logistics_fit": 90, "ghost_score": 5})
    return rt, llm, job_id


def wait(future) -> None:  # noqa: ANN001
    future.result(timeout=60)


def test_full_workflow_with_validation_loop_edits_restart_and_apply(app_rt, cfg, engine):
    rt, llm, job_id = app_rt
    wait(triage.shortlist(rt, job_id, reason="great fit", emphasis="API platforms"))
    assert queries.get_job(engine, job_id)["status"] == "in_workflow"
    assert queries.get_job(engine, job_id)["thread_id"] == f"job-{job_id}"
    v = workflow.view(rt, job_id)
    assert v["gate"]["gate"] == "review_resume" and v["gate"]["errors"] == []
    assert v["state"]["coverage"] and v["state"]["keyword_coverage"]["matched"]
    first_doc = v["gate"]["document_id"]
    assert first_doc and queries.get_document(engine, first_doc)["kind"] == "resume"
    assert [d["decision"] if "decision" in d else d["action"] for d in queries.decisions_for_job(engine, job_id)] == [
        "shortlist"]

    # 1) A fabricated metric hand-typed into the editor is caught and nothing is saved.
    docs_before = len(queries.documents_for_job(engine, job_id))
    tampered = v["editor"].replace("onboarding time 60%", "onboarding time 85%")
    wait(workflow.resume(rt, job_id, workflow.edit_payload("resume", tampered)))
    v = workflow.view(rt, job_id)
    assert v["gate"]["gate"] == "review_resume"
    assert any(e["problem"] == "unsupported_number" and "85%" in e["detail"] for e in v["gate"]["errors"])
    assert len(queries.documents_for_job(engine, job_id)) == docs_before
    # accept is refused while errors stand
    wait(workflow.resume(rt, job_id, {"action": "accept"}))
    assert "cannot be accepted" in workflow.view(rt, job_id)["gate"]["message"]

    # 2) Revise: the model fabricates three times -> loop exhausts max_draft_attempts -> interrupt with errors.
    llm.resumes = [FABRICATED_SUMMARY, FABRICATED_SUMMARY, FABRICATED_SUMMARY, GOOD_RESUME]
    before = llm.calls["draft_resume"]
    wait(workflow.resume(rt, job_id, {"action": "revise", "instructions": "tighten bullets"}))
    v = workflow.view(rt, job_id)
    assert llm.calls["draft_resume"] - before == cfg.tailoring.max_draft_attempts
    assert v["gate"]["gate"] == "review_resume" and v["gate"]["errors"]
    assert "failed validation 3 times" in v["gate"]["message"]

    # 3) Revise again: a clean draft passes and can be accepted.
    wait(workflow.resume(rt, job_id, {"action": "revise", "instructions": "use only the facts"}))
    v = workflow.view(rt, job_id)
    assert v["gate"]["errors"] == [] and v["gate"]["document_id"]
    wait(workflow.resume(rt, job_id, {"action": "accept"}))
    v = workflow.view(rt, job_id)
    assert v["gate"]["gate"] == "review_cover" and v["gate"]["document_id"]
    wait(workflow.resume(rt, job_id, {"action": "accept"}))
    v = workflow.view(rt, job_id)
    assert v["gate"]["gate"] == "approve_and_apply", v["failure"]
    assert queries.get_job(engine, job_id)["status"] == "ready"
    paths = [Path(p) for p in v["state"]["rendered_paths"]]
    assert len(paths) == 2 and all(p.exists() and len(str(p)) < 240 for p in paths)
    folder = paths[0].parent.name
    assert folder.endswith(f"-{job_id}") and ":" not in folder and "/" not in folder
    assert render.check_docx_ats(paths[0]) == []
    assert any("Upload" in s for s in v["gate"]["checklist"]["steps"])

    # 4) Full restart mid-interrupt: new runtime, new checkpointer connection, same thread resumes.
    rt2 = init_runtime(cfg, llm_factory=lambda: llm, embedder_factory=rt.embedder_factory)
    assert workflow.pending_gate(rt2, job_id)["gate"] == "approve_and_apply"
    wait(workflow.resume(rt2, job_id, {"action": "confirm", "method": "deep_link", "confirmation_ref": "R-1",
                                       "notes": "applied via Workday"}))
    assert queries.get_job(engine, job_id)["status"] == "applied"
    apps = queries.list_applications(engine)
    assert len(apps) == 1 and apps[0]["confirmation_ref"] == "R-1" and len(apps[0]["document_ids"]) == 2
    assert apps[0]["follow_up_at"] > apps[0]["submitted_at"]
    assert workflow.view(rt2, job_id)["done"]


def test_a_fixable_draft_is_repaired_instead_of_redrafted(app_rt, engine):
    """A bullet with a number its fact lacks goes back to the fact's own words: one draft, not two, and the review
    page says what was changed."""
    rt, llm, job_id = app_rt
    llm.resumes = [FABRICATED, GOOD_RESUME]
    wait(triage.shortlist(rt, job_id))
    v = workflow.view(rt, job_id)
    assert llm.calls["draft_resume"] == 1
    assert v["gate"]["gate"] == "review_resume" and v["gate"]["errors"] == [] and v["gate"]["document_id"]
    assert "Fixed automatically" in v["gate"]["message"] and "ach-example-gateway" in v["gate"]["message"]
    saved = queries.get_document(engine, v["gate"]["document_id"])["markdown"]
    assert "85%" not in saved and "60%" in saved


def test_tailoring_checks_the_posting_is_still_up(app_rt, engine, monkeypatch):
    """A posting the board no longer lists isn't tailored: the job is marked closed, and "Tailor anyway" still
    works when the check was wrong. An unclear answer lets tailoring go ahead."""
    from jobhunt.sources import liveness

    rt, llm, job_id = app_rt
    monkeypatch.setattr(liveness, "check", lambda job, client=None: liveness.Liveness(
        "closed", "Workday no longer lists this posting"))
    with pytest.raises(triage.PostingClosed, match="looks closed"):
        triage.shortlist(rt, job_id, verify=True)
    job = queries.get_job(engine, job_id)
    assert job["status"] == "closed" and job["status_reason"].startswith("posting closed")
    assert llm.calls.get("draft_resume", 0) == 0                          # nothing was drafted

    wait(triage.shortlist(rt, job_id, reason="tailored anyway", verify=True))   # closed: no second check
    assert queries.get_job(engine, job_id)["status"] == "in_workflow"


def test_an_unclear_link_check_does_not_block_tailoring(app_rt, engine, monkeypatch):
    from jobhunt.sources import liveness

    rt, _, job_id = app_rt
    monkeypatch.setattr(liveness, "check", lambda job, client=None: liveness.Liveness("unknown", "timed out"))
    wait(triage.shortlist(rt, job_id, verify=True))
    assert queries.get_job(engine, job_id)["status"] == "in_workflow"


def test_ats_super_match_asks_before_adding_keywords_and_a_headline(app_rt, cfg, engine):
    """[Amendment A4] The workflow stops before drafting with the posting keywords your facts lack and a headline;
    only the ticked ones are used. Ticked keywords join the fact bank, so the fact check still holds."""
    from docx import Document

    from jobhunt.facts import bank as fact_bank

    rt, llm, job_id = app_rt
    llm.requirements_keywords = ["waterfall", "budget management"]       # not in the example facts
    wait(triage.shortlist(rt, job_id, generation_mode="ats_super_match"))
    v = workflow.view(rt, job_id)
    assert v["gate"]["gate"] == "review_ats" and llm.calls.get("draft_resume", 0) == 0
    offered = [k["keyword"] for k in v["gate"]["keywords"]]
    assert "waterfall" in offered                       # no fact mentions it
    assert "budget management" not in offered          # already supported: "Managed a $2.5M ... budget"
    assert v["gate"]["headline"].startswith("Senior Architect")                   # "Sr." spelled out, no location
    pick = "waterfall"

    wait(workflow.resume(rt, job_id, {"action": "apply", "keywords": [pick, "not offered"], "use_headline": True,
                                      "headline": "Senior Integration Architect"}))
    v = workflow.view(rt, job_id)
    assert v["gate"]["gate"] == "review_resume" and llm.calls["draft_resume"] >= 1
    assert v["state"]["ats_confirmed_keywords"] == [pick] and v["state"]["generation_mode"] == "ats_super_match"
    # for this job only: a source in the workflow, nothing in facts.yaml or the fact table
    assert any(f["id"] == "ats-confirmed-keywords" and pick in f["tags"] for f in v["state"]["candidate_facts"])
    assert not any(pick in (f.tags or []) for f in fact_bank.load_yaml(cfg.paths.facts).facts)
    assert not any("ats-confirmed" in f["id"] for f in queries.list_facts(engine))

    wait(workflow.resume(rt, job_id, {"action": "accept"}))
    wait(workflow.resume(rt, job_id, {"action": "accept"}))
    paths = workflow.view(rt, job_id)["state"]["rendered_paths"]
    texts = [p.text for p in Document(paths[0]).paragraphs]
    assert "Senior Integration Architect" in texts[:4]                                # under the name


def test_automated_mode_drafts_straight_away(app_rt):
    rt, llm, job_id = app_rt
    wait(triage.shortlist(rt, job_id, generation_mode="automated"))
    assert workflow.view(rt, job_id)["gate"]["gate"] == "review_resume"


def test_invalid_action_for_gate_is_refused(app_rt):
    rt, _, job_id = app_rt
    wait(triage.shortlist(rt, job_id))
    with pytest.raises(workflow.WorkflowError):
        workflow.resume(rt, job_id, {"action": "confirm"})


def test_abandon(app_rt, engine):
    rt, _, job_id = app_rt
    wait(triage.shortlist(rt, job_id))
    wait(workflow.resume(rt, job_id, {"action": "abandon"}))
    assert queries.get_job(engine, job_id)["status"] == "abandoned"
    assert workflow.view(rt, job_id)["done"]


def test_reject_requires_reason_and_records_decision(app_rt, engine):
    rt, _, job_id = app_rt
    with pytest.raises(triage.TriageError):
        triage.reject(rt, job_id, "  ")
    triage.reject(rt, job_id, "too junior")
    assert queries.get_job(engine, job_id)["status"] == "rejected"
    d = queries.decisions_for_job(engine, job_id)[0]
    assert d["action"] == "reject" and d["reason"] == "too junior" and d["total_at_decision"] == 82.0


def test_retrieval_builds_skeleton_plus_achievements(app_rt, engine, hashing_embedder):
    candidates, coverage = bank.retrieve(engine, hashing_embedder,
                                         [{"text": "MuleSoft API gateway", "kind": "must", "keywords": ["mulesoft"]}],
                                         top_k=2, threshold=0.6)
    kinds = [c["kind"] for c in candidates]
    assert kinds.count("achievement") == 2
    assert {"role", "summary", "skill", "education", "cert"} <= set(kinds)
    assert coverage[0]["status"] in ("covered", "partial", "gap")


REQ_DEGREE = "Bachelor's degree or equivalent experience (HS diploma + 4 years relevant experience)"
EDUCATION = [{"id": "edu-bs", "kind": "education", "title": "Bachelor of Science (BS) Computer Science",
              "text": "Binghamton University | Bachelor of Science (BS) Computer Science, Magna Cum Laude"},
             {"id": "edu-ms", "kind": "education", "title": "Master of Science (MS) Information Systems",
              "text": "Drexel University | Master of Science (MS) Information Systems"}]


@pytest.mark.parametrize(("requirement", "met_by"), [
    (REQ_DEGREE, ["edu-bs", "edu-ms"]),                                   # the owner's case: this is not a gap
    ("BS in Computer Science or related field", ["edu-bs", "edu-ms"]),    # a higher degree meets a lower one
    ("Master's degree in Information Systems preferred", ["edu-ms"]),
    ("PhD in Machine Learning", []),                                      # a real gap
    ("BS/MS in Computer Science", ["edu-bs", "edu-ms"]),                  # either degree
    ("Master's or PhD in Computer Science", ["edu-ms"]),
    ("Proficiency in MS Office", None),                                   # not a degree
    ("8+ years leading integration teams", None),                         # not about a degree: left to matching
])
def test_a_degree_requirement_is_decided_by_the_education_facts(requirement, met_by):
    assert bank.education_covers(requirement, EDUCATION) == met_by


def test_a_degree_requirement_is_covered_when_the_facts_hold_the_degree(app_rt, engine, hashing_embedder):
    _, coverage = bank.retrieve(engine, hashing_embedder,
                                [{"text": REQ_DEGREE, "kind": "must", "keywords": ["bachelor's degree"]}],
                                top_k=2, threshold=0.6)
    assert coverage[0]["status"] == "covered"
    assert coverage[0]["covered_by"] and all(f.startswith("edu") for f in coverage[0]["covered_by"])


def test_template_that_is_a_finished_resume_is_refused(tmp_path):
    # A resume saved from Word usually has no "List Bullet" style and uses per-line formatting, so it can't act as a
    # template: bullets would silently lose their markers. Rendering must fail loudly instead.
    doc = Document()
    doc.add_paragraph("JANE PUBLIC").runs[0].bold = True
    styles = doc.styles.element
    for st in list(styles):
        if st.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}styleId") == "ListBullet":
            styles.remove(st)
    finished = tmp_path / "my-resume.docx"
    doc.save(str(finished))
    with pytest.raises(render.RenderError, match="missing Word styles List Bullet"):
        render.markdown_to_docx(GOOD_RESUME, tmp_path / "out.docx", template=finished, name="Jane", contact_line="",
                                kind="resume")


def test_default_template_is_styles_only_and_black(tmp_path):
    path = render.make_default_template(tmp_path / "t.docx", margin_in=0.5)
    doc = Document(str(path))
    assert not any(p.text.strip() for p in doc.paragraphs)
    assert render.missing_styles(doc) == [] and render.check_docx_ats(path) == []
    h1 = doc.styles["Heading 1"]
    assert h1.font.name == "Calibri" and h1.font.color.rgb == (0, 0, 0) and h1.font.bold
    assert round(doc.sections[0].left_margin.inches, 2) == 0.5


def test_role_headings_are_split_for_resume_parsers(tmp_path, cfg):
    """Workday and other parsers segment work history out of the normal text flow: employer, title and dates each
    need their own plain line. A heading *style* on the employer reads as a section break, and the parser then takes
    the employer from bullet text instead."""
    template = render.make_default_template(tmp_path / "template.docx")
    md = GOOD_RESUME.replace("### Lead Integration Architect | Example Insurance Co | Mar 2019 – Present",
                             "### Solution Lead – Enterprise Integration | Example Insurance Co | Mar 2019 – Present")
    out = render.markdown_to_docx(md, tmp_path / "split.docx", template=template, name="Jane", contact_line="x",
                                  kind="resume", role_layout="separate_lines")
    doc = Document(str(out))
    lines = [(p.style.name, p.text) for p in doc.paragraphs if p.text.strip()]
    start = next(i for i, (_, t) in enumerate(lines) if t == "WORK EXPERIENCE")
    assert lines[start + 1] == ("Normal", "Example Insurance Co")             # employer alone, body text not a heading
    assert lines[start + 2] == ("Normal", "Solution Lead, Enterprise Integration")  # a dash would split the title
    assert lines[start + 3] == ("Normal", "03/2019 - Present")                # MM/YYYY, plain hyphen
    assert doc.paragraphs[lines_index(doc, "Example Insurance Co")].runs[0].bold  # bold keeps it readable
    edu = next(i for i, (_, t) in enumerate(lines) if t == "EDUCATION")
    assert lines[edu + 1] == ("Normal", "Example University")
    assert render.check_docx_ats(out) == []

    compact = render.markdown_to_docx(md, tmp_path / "one.docx", template=template, name="Jane", contact_line="x",
                                      kind="resume", role_layout="one_line")
    headings = [p.text for p in Document(str(compact)).paragraphs if p.style.name == "Heading 2"]
    assert any("|" in h and "Example Insurance Co" in h for h in headings)


def lines_index(doc, text: str) -> int:
    return next(i for i, p in enumerate(doc.paragraphs) if p.text == text)


@pytest.mark.parametrize(("title", "expected"), [
    ("Sr. IT Program Manager", "Senior IT Program Manager"),
    ("Director, Tech Lead - PGIM Fixed Income (Hybrid - Newark, NJ)", "Director, Tech Lead - PGIM Fixed Income"),
    ("Integration Lead - Remote", "Integration Lead"),
    ("Senior Manager, IT Delivery R12345", "Senior Manager, IT Delivery"),
])
def test_the_ats_headline_is_the_title_without_codes_or_locations(title, expected):
    from jobhunt.graphs.application.nodes import ats_headline

    assert ats_headline(title, []) == expected
    assert ats_headline(title, [{"keyword": "mulesoft", "kind": "must"}]) == f"{expected} | Mulesoft"


def test_skills_groups_have_a_bold_label_and_stay_one_plain_line_each(tmp_path):
    """Readable for a person (the group label stands out), unchanged for a parser (one line of text per group)."""
    md = GOOD_RESUME.replace("Integration: MuleSoft, REST APIs, API gateways, Kafka, AWS",
                             "Integration: MuleSoft, REST APIs, API gateways\nCloud: AWS, Kafka")
    out = render.markdown_to_docx(md, tmp_path / "r.docx", template=render.make_default_template(tmp_path / "t.docx"),
                                  name="Jane Public", contact_line="NJ", kind="resume")
    assert render.check_docx_ats(out) == []
    paras = Document(str(out)).paragraphs
    i = next(k for k, p in enumerate(paras) if p.text == "SKILLS")
    first, second = paras[i + 1], paras[i + 2]
    assert first.text == "Integration: MuleSoft, REST APIs, API gateways" and second.text == "Cloud: AWS, Kafka"
    assert first.runs[0].text == "Integration:" and first.runs[0].bold and not first.runs[1].bold


def test_a_rendered_document_opens_in_word_only_from_the_output_folder(app_rt, cfg, engine, monkeypatch):
    import os
    import sys

    rt, _, job_id = app_rt
    opened = []
    monkeypatch.setattr(sys, "platform", "win32")
    monkeypatch.setattr(os, "startfile", lambda path: opened.append(path), raising=False)
    inside = cfg.paths.output_root / "acme" / "Jane-Resume.docx"
    inside.parent.mkdir(parents=True, exist_ok=True)
    inside.write_bytes(b"docx")
    doc_id = queries.insert_document(engine, job_id=job_id, kind="resume", markdown=GOOD_RESUME, fact_map=[],
                                     fact_ids=[], coverage=None)
    queries.set_document_files(engine, doc_id, [str(inside)])
    assert workflow.open_document(rt, doc_id) == inside.resolve() and opened == [str(inside.resolve())]

    outside = cfg.paths.output_root.parent / "elsewhere.docx"
    outside.write_bytes(b"docx")
    queries.set_document_files(engine, doc_id, [str(outside)])
    with pytest.raises(workflow.WorkflowError):
        workflow.open_document(rt, doc_id)
    assert len(opened) == 1

