from __future__ import annotations

import os

import numpy as np
import pytest
from openpyxl import load_workbook

from jobhunt.db import export, locks, queries, vectors
from jobhunt.db.engine import age_days, parse_iso, to_iso, utcnow_iso
from jobhunt.sources.normalize import normalize_raw
from tests.conftest import raw_job


def test_pragmas_applied_on_every_connection(engine):
    p = queries.pragma_values(engine)
    assert p["journal_mode"].lower() == "wal"
    assert p["busy_timeout"] == 10000
    assert p["foreign_keys"] == 1
    assert p["synchronous"] == 1  # NORMAL


def test_timestamp_format_is_canonical():
    now = utcnow_iso()
    assert now.endswith("Z") and len(now) == 20
    assert to_iso(parse_iso(now)) == now
    assert parse_iso("2026-09-14T13:45:02.123+02:00") is not None
    assert age_days(now) < 0.01


UNICODE_DESC = (
    "<p>Senior Architect — “Integration” lead at Zürich Re. Salary range: $180,000–$210,000. "
    "Team: Платформа / 東京 office. Don’t miss it…</p>"
)


def test_unicode_posting_round_trips_intact(engine, tmp_path):
    rec = normalize_raw(raw_job(1, title="Architect — “API” Platform 東京", description=UNICODE_DESC))
    outcome, job_id = queries.upsert_job(engine, rec)
    assert outcome == "new"
    job = queries.get_job(engine, job_id)
    assert job["title"] == "Architect — “API” Platform 東京"
    for fragment in ("—", "“Integration”", "Zürich", "Платформа", "東京", "Don’t", "…"):
        assert fragment in job["description_text"]
    # FTS finds non-ASCII text too
    assert job_id in [i for i, _ in queries.fts_search(engine, "Zürich")]
    # and the xlsx export carries it unchanged
    path = export.write_xlsx([job], tmp_path / "u.xlsx")
    ws = load_workbook(path).active
    cells = [c.value for c in ws[2]]
    assert "Architect — “API” Platform 東京" in cells


def test_upsert_seen_changed_and_sighting(engine):
    rec = normalize_raw(raw_job(2))
    assert queries.upsert_job(engine, rec)[0] == "new"
    assert queries.upsert_job(engine, rec)[0] == "seen"
    changed = normalize_raw(raw_job(2, description="<p>totally different text</p>"))
    outcome, job_id = queries.upsert_job(engine, changed)
    assert outcome == "changed"
    assert queries.upsert_job(engine, {**changed, "content_hash": "x"}, sighting_only=True)[0] == "seen"
    job = queries.get_job(engine, job_id)
    assert job["seen_count"] == 4
    assert job["description_text"] == "totally different text"
    assert queries.upsert_job(engine, normalize_raw(raw_job(99)), sighting_only=True) == ("unknown_sighting", None)


def test_status_transitions_are_validated(engine):
    _, job_id = queries.upsert_job(engine, normalize_raw(raw_job(3)))
    with pytest.raises(queries.InvalidTransition):
        queries.set_job_status(engine, job_id, "rejected")  # reject needs a score first
    queries.set_job_status(engine, job_id, "scored")
    queries.set_job_status(engine, job_id, "snoozed", snooze_until="2000-01-01T00:00:00Z")
    assert queries.wake_snoozed(engine) == 1
    assert queries.get_job(engine, job_id)["status"] == "scored"
    with pytest.raises(queries.InvalidTransition):
        queries.set_job_status(engine, job_id, "discovered")


def test_fts_triggers_follow_updates(engine):
    _, job_id = queries.upsert_job(engine, normalize_raw(raw_job(4, description="<p>kafka streaming</p>")))
    assert [i for i, _ in queries.fts_search(engine, "kafka")] == [job_id]
    queries.upsert_job(engine, normalize_raw(raw_job(4, description="<p>rabbitmq queues</p>")))
    assert queries.fts_search(engine, "kafka") == []
    assert [i for i, _ in queries.fts_search(engine, "rabbitmq")] == [job_id]


def test_lock_refuses_second_holder_and_reclaims_stale(engine):
    locks.acquire(engine, "discovery_run", note="first")
    with pytest.raises(locks.LockHeld):
        locks.acquire(engine, "discovery_run", note="second")
    # A stale lock (older than the threshold) is reclaimed.
    with engine.begin() as c:
        c.exec_driver_sql("UPDATE locks SET acquired_at = '2000-01-01T00:00:00Z' WHERE name = 'discovery_run'")
    locks.acquire(engine, "discovery_run", note="third", stale_minutes=30)
    assert queries.lock_get(engine, "discovery_run")["note"] == "third"
    locks.release(engine, "discovery_run")
    assert queries.lock_get(engine, "discovery_run") is None


def test_lock_held_by_live_other_process_is_not_reclaimed(engine):
    parent = os.getppid()
    assert locks.pid_alive(parent)
    with engine.begin() as c:
        c.exec_driver_sql("INSERT INTO locks (name, acquired_at, pid, note) VALUES ('discovery_run', ?, ?, 'cli')",
                          (utcnow_iso(), parent))
    with pytest.raises(locks.LockHeld):
        locks.acquire(engine, "discovery_run")


def test_lock_held_by_dead_process_is_reclaimed(engine):
    with engine.begin() as c:
        c.exec_driver_sql("INSERT INTO locks (name, acquired_at, pid, note) VALUES ('thread:job-1', ?, 999999, 'x')",
                          (utcnow_iso(),))
    locks.acquire(engine, "thread:job-1")
    assert queries.lock_get(engine, "thread:job-1")["pid"] == os.getpid()


def test_vectors_roundtrip_and_topk():
    v = vectors.to_blob([3.0, 4.0] + [0.0] * 382)
    assert len(v) == 1536
    arr = vectors.from_blob(v)
    assert np.isclose(np.linalg.norm(arr), 1.0)
    rows = [{"id": i, "embedding": vectors.to_blob(np.eye(384)[i])} for i in range(5)]
    ids, mat = vectors.matrix(rows)
    best = vectors.top_k(mat, np.eye(384)[3], k=2)
    assert ids[best[0][0]] == 3 and np.isclose(best[0][1], 1.0)


def test_backup_uses_vacuum_into(engine, tmp_path):
    queries.upsert_job(engine, normalize_raw(raw_job(5)))
    out = tmp_path / "exports" / "backup.db"
    queries.vacuum_into(engine, out)
    assert out.exists() and out.stat().st_size > 0
    with pytest.raises(FileExistsError):
        queries.vacuum_into(engine, out)


def test_export_truncates_long_descriptions_only_in_export(engine, tmp_path):
    long_desc = "x" * 19_000
    _, job_id = queries.upsert_job(engine, normalize_raw(raw_job(6, description=long_desc)))
    row = queries.get_job(engine, job_id)
    row["description_text"] = "y" * 40_000  # simulate an over-long value
    ws = load_workbook(export.write_xlsx([row], tmp_path / "t.xlsx")).active
    desc = [c.value for c in ws[2]][-1]
    assert len(desc) == 32_001 and desc.endswith("…")
    assert len(queries.get_job(engine, job_id)["description_text"]) == 19_000
