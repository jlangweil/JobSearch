"""One-way xlsx export (SPEC 4.17). Nothing reads a spreadsheet back."""

from __future__ import annotations

import io
from collections.abc import Sequence
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

EXCEL_CELL_LIMIT = 32_000

COLUMNS: list[tuple[str, str, int]] = [
    ("id", "Job ID", 8),
    ("total", "Score", 8),
    ("ghost_score", "Ghost", 7),
    ("company_name_raw", "Company", 22),
    ("title", "Title", 40),
    ("location_raw", "Location", 24),
    ("remote_type", "Remote", 10),
    ("employment_type", "Employment", 11),
    ("comp_min", "Comp min", 11),
    ("comp_max", "Comp max", 11),
    ("posted_at", "Posted (UTC)", 20),
    ("status", "Status", 12),
    ("source", "Source", 12),
    ("skill_match", "Skill", 7),
    ("role_fit", "Role fit", 8),
    ("domain_fit", "Domain", 7),
    ("seniority_fit", "Seniority", 9),
    ("comp_fit", "Comp fit", 8),
    ("logistics_fit", "Logistics", 9),
    ("rationale", "Rationale", 60),
    ("url", "URL", 40),
    ("apply_url", "Apply URL", 40),
    ("description_text", "Description", 80),
]


def _cell(value):  # noqa: ANN001, ANN202
    if isinstance(value, str) and len(value) > EXCEL_CELL_LIMIT:
        return value[:EXCEL_CELL_LIMIT] + "…"
    return value


def build_workbook(rows: Sequence[dict], sheet_title: str = "jobs") -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title[:31]
    ws.append([label for _, label, _ in COLUMNS])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for r in rows:
        ws.append([_cell(r.get(key)) for key, _, _ in COLUMNS])
    for i, (_, _, width) in enumerate(COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "A2"
    return wb


def write_xlsx(rows: Sequence[dict], out_path: Path) -> Path:
    out_path = Path(out_path).resolve()
    out_path.parent.mkdir(parents=True, exist_ok=True)
    build_workbook(rows).save(out_path)
    return out_path


def xlsx_bytes(rows: Sequence[dict]) -> bytes:
    buf = io.BytesIO()
    build_workbook(rows).save(buf)
    return buf.getvalue()
